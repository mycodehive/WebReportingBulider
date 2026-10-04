"""Read-only data connectors and a bounded, portable query engine.

External connections are implementation-ready, not integration-certified. Drivers
are imported only when selected. SQL identifiers always come from introspection;
report definitions never contain executable SQL, Python, or endpoint templates.
"""
from __future__ import annotations

import base64
import csv
import datetime as dt
import functools
import http.client
import ipaddress
import io
import json
import math
import operator
import re
import shutil
import socket
import sqlite3
import ssl
import subprocess
from contextlib import contextmanager
from decimal import Decimal, InvalidOperation
from pathlib import Path
from urllib.parse import quote, urlsplit, urlencode


class DataError(Exception):
    def __init__(self, code, message):
        self.code = code
        self.message = message
        super().__init__(message)


_CATALOG = [
    ("excel", "Excel (.xlsx/.xlsm/.xls)", "file"),
    ("csv", "CSV / TSV", "file"),
    ("sqlite", "SQLite", "file"),
    ("postgresql", "PostgreSQL", "database"),
    ("mariadb", "MariaDB / MySQL", "database"),
    ("oracle", "Oracle", "database"),
    ("mssql", "Microsoft SQL Server", "database"),
    ("google_sheets", "Google Sheets", "api"),
    ("access", "Access (MDBTools, read-only)", "file"),
    ("rest", "REST JSON (fixed HTTPS endpoint)", "api"),
]
_ALIASES = {"xlsx": "excel", "xls": "excel", "mysql": "mariadb", "sheets": "google_sheets", "sqlserver": "mssql"}
_PLUGIN_CONNECTORS = {}
_BASE_CAPABILITIES = {
    "read_only": True, "schema": True, "filters": True, "sorts": True,
    "groups": True, "aggregates": True, "joins": False, "raw_sql": False,
    "query_pushdown": False, "bounded_snapshot": True,
}
_SOURCE_BYTES = 20 * 1024 * 1024
_MAX_QUERY_ROWS = 100000


def register_connector(kind, connector):
    """Server administrator extension point; imported projects cannot install code.

    connector must expose introspect(config), read(config, object_mapping,
    columns, limit). read must return dictionaries and obey the limit. Plugins
    are trusted server code, installed/reviewed by the administrator.
    """
    if not re.fullmatch(r"[a-z][a-z0-9_]{1,63}", kind) or kind in {x[0] for x in _CATALOG}:
        raise DataError("INVALID_CONNECTOR", "커넥터 식별자가 올바르지 않거나 이미 등록되어 있습니다.")
    if not callable(getattr(connector, "introspect", None)) or not callable(getattr(connector, "read", None)):
        raise DataError("INVALID_CONNECTOR", "커넥터 인터페이스가 올바르지 않습니다.")
    _PLUGIN_CONNECTORS[kind] = connector


def connector_catalog():
    result = [{"kind": k, "name": n, "label": n, "category": c,
               "verified": False, "verification": "external_connection_not_verified",
               "capabilities": dict(_BASE_CAPABILITIES)} for k, n, c in _CATALOG]
    result.extend({"kind": k, "name": k, "label": k, "category": "plugin", "verified": False,
                   "capabilities": dict(_BASE_CAPABILITIES)} for k in _PLUGIN_CONNECTORS)
    return result


def _kind(kind):
    kind = _ALIASES.get(kind, kind)
    if kind not in {x[0] for x in _CATALOG} and kind not in _PLUGIN_CONNECTORS:
        raise DataError("UNSUPPORTED_CONNECTOR", "지원하지 않는 데이터 커넥터입니다.")
    return kind


def _integer(value, default, low, high):
    try:
        number = int(default if value is None else value)
    except (TypeError, ValueError):
        raise DataError("INVALID_CONFIG", "숫자 설정이 올바르지 않습니다.") from None
    if not low <= number <= high:
        raise DataError("INVALID_CONFIG", "숫자 설정이 허용 범위를 벗어났습니다.")
    return number


def _file(config, suffixes=None):
    try:
        path = Path(config["path"])
    except (KeyError, TypeError):
        raise DataError("MISSING_FILE", "서버에 등록된 자료 파일이 필요합니다.") from None
    if not path.is_absolute() or not path.is_file() or path.is_symlink():
        raise DataError("MISSING_FILE", "서버에 등록된 자료 파일을 찾을 수 없습니다.")
    if suffixes and path.suffix.lower() not in suffixes:
        raise DataError("INVALID_FILE", "지원하지 않는 자료 파일 형식입니다.")
    if path.stat().st_size > _integer(config.get("max_file_bytes"), 50 * 1024 * 1024, 1, 512 * 1024 * 1024):
        raise DataError("DATA_LIMIT", "자료 파일이 허용 크기를 초과했습니다.")
    return path


def _headers(values):
    """Retain physical labels; create deterministic identifiers for blank/duplicate headers."""
    result, seen = [], set()
    for idx, item in enumerate(values):
        base = str(item).strip() if item is not None and str(item).strip() else f"column_{idx + 1}"
        name, suffix = base, 2
        while name in seen:
            name, suffix = f"{base}_{suffix}", suffix + 1
        seen.add(name)
        result.append(name)
    return result


def _value_type(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, Decimal):
        return "decimal"
    if isinstance(value, float):
        return "float"
    if isinstance(value, dt.datetime):
        return "datetime"
    if isinstance(value, dt.date):
        return "date"
    if isinstance(value, (bytes, bytearray, memoryview)):
        return "binary"
    if isinstance(value, (dict, list)):
        return "json"
    return "string"


def _columns(names, rows, *, inferred=True):
    columns = []
    for name in names:
        values = [row.get(name) for row in rows]
        types = {_value_type(v) for v in values if v is not None}
        if types <= {"integer", "float", "decimal"} and types:
            typ = "decimal" if "decimal" in types else "float" if "float" in types else "integer"
        else:
            typ = next(iter(types)) if len(types) == 1 else "string" if types else "unknown"
        columns.append({"column": name, "name": name, "label": name, "type": typ,
                        # A sample without NULL cannot prove a source invariant.
                        "nullable": True if inferred else not values or any(v is None for v in values),
                        "sample_contains_null": not values or any(v is None for v in values),
                        "inferred": inferred, "sample_rows": len(rows)})
    return columns


def _object(name, columns, schema=None, typ="table", **extra):
    return {"object": name, "name": name, "schema": schema, "type": typ, "columns": columns, **extra}


@contextmanager
def _sqlite(config):
    path = _file(config, {".sqlite", ".sqlite3", ".db"})
    conn = sqlite3.connect(path.as_uri() + "?mode=ro", uri=True, timeout=5)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA query_only=ON")
    ticks = [0]
    def progress():
        ticks[0] += 1
        return int(ticks[0] > 10000)
    conn.set_progress_handler(progress, 1000)
    try:
        yield conn
    finally:
        conn.close()


def _identifier(name):
    # This is quoting, never authorization. Callers must verify against metadata.
    if not isinstance(name, str) or not name or "\x00" in name:
        raise DataError("INVALID_MAPPING", "자료 식별자가 올바르지 않습니다.")
    return '"' + name.replace('"', '""') + '"'


def _sql_type(name):
    text = str(name).upper()
    if any(x in text for x in ("INT", "SERIAL")):
        return "integer"
    if any(x in text for x in ("DECIMAL", "NUMERIC", "NUMBER", "MONEY")):
        return "decimal"
    if any(x in text for x in ("FLOAT", "REAL", "DOUBLE")):
        return "float"
    if "BOOL" in text or text == "BIT":
        return "boolean"
    if "TIMESTAMP" in text or "DATETIME" in text:
        return "datetime"
    if text == "DATE":
        return "date"
    if "BLOB" in text or "BINARY" in text or "BYTEA" in text:
        return "binary"
    if "JSON" in text:
        return "json"
    return "string"


def _sqlite_schema(config):
    objects = []
    with _sqlite(config) as conn:
        for row in conn.execute("SELECT name,type FROM sqlite_master WHERE type IN ('table','view') AND name NOT LIKE 'sqlite_%' ORDER BY name"):
            columns = []
            for col in conn.execute("PRAGMA table_info(" + _identifier(row["name"]) + ")"):
                columns.append({"column": col["name"], "name": col["name"], "type": _sql_type(col["type"]),
                                "native_type": col["type"], "nullable": not bool(col["notnull"] or col["pk"]),
                                "primary_key": bool(col["pk"]), "inferred": False})
            objects.append(_object(row["name"], columns, typ=row["type"]))
    return objects


def _csv_records(config, limit):
    path = _file(config, {".csv", ".tsv", ".txt"})
    delimiter = config.get("delimiter", "\t" if path.suffix.lower() == ".tsv" else ",")
    if not isinstance(delimiter, str) or len(delimiter) != 1:
        raise DataError("INVALID_CONFIG", "CSV 구분자는 한 글자여야 합니다.")
    encoding = config.get("encoding", "utf-8-sig")
    if encoding not in {"utf-8", "utf-8-sig", "cp949", "euc-kr", "utf-16"}:
        raise DataError("INVALID_CONFIG", "지원하지 않는 문자 인코딩입니다.")
    header_row = _integer(config.get("header_row"), 1, 1, 10000)
    with path.open("r", encoding=encoding, newline="") as stream:
        reader = csv.reader(stream, delimiter=delimiter)
        for _ in range(header_row - 1):
            next(reader, None)
        names = _headers(next(reader, []))
        rows = []
        for values in reader:
            if len(values) > len(names):
                raise DataError("INVALID_FILE", "CSV 데이터의 열 수가 헤더보다 많습니다.")
            if not any(str(v).strip() for v in values):
                continue
            rows.append({name: values[i] if i < len(values) and values[i] != "" else None for i, name in enumerate(names)})
            if len(rows) >= limit:
                break
    return names, rows


def _excel_records(config, limit, selected=None):
    path = _file(config, {".xlsx", ".xlsm", ".xls"})
    header_row = _integer(config.get("header_row"), 1, 1, 10000)
    if path.suffix.lower() == ".xls":
        try:
            import xlrd
        except ImportError:
            raise DataError("DRIVER_MISSING", "구형 Excel(.xls)을 읽으려면 xlrd를 설치하세요.") from None
        book = xlrd.open_workbook(str(path), on_demand=True)
        try:
            names = book.sheet_names()
            results = []
            for sheet_name in names if selected is None else [selected]:
                if sheet_name not in names:
                    raise DataError("INVALID_MAPPING", "Excel 시트를 찾을 수 없습니다.")
                sheet = book.sheet_by_name(sheet_name)
                def value(cell):
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        return xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)
                    if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
                        return None
                    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        return bool(cell.value)
                    return cell.value
                headers = _headers([value(c) for c in sheet.row(header_row - 1)]) if sheet.nrows >= header_row else []
                records = []
                for idx in range(header_row, sheet.nrows):
                    vals = [value(c) for c in sheet.row(idx)]
                    if not any(v is not None for v in vals):
                        continue
                    records.append(dict(zip(headers, vals)))
                    if len(records) >= limit:
                        break
                results.append((sheet_name, headers, records))
            return results
        finally:
            book.release_resources()
    try:
        import openpyxl
    except ImportError:
        raise DataError("DRIVER_MISSING", "Excel 커넥터에 openpyxl을 설치하세요.") from None
    import zipfile
    with zipfile.ZipFile(path) as archive:
        if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
            raise DataError("DATA_LIMIT", "압축 해제된 Excel 자료가 허용 크기를 초과했습니다.")
    book = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        results = []
        for sheet_name in book.sheetnames if selected is None else [selected]:
            if sheet_name not in book.sheetnames:
                raise DataError("INVALID_MAPPING", "Excel 시트를 찾을 수 없습니다.")
            sheet = book[sheet_name]
            if sheet.max_column is not None and sheet.max_column > 4096:
                raise DataError("DATA_LIMIT", "Excel 열 수가 허용 범위를 초과했습니다.")
            iterator = sheet.iter_rows(min_row=header_row, values_only=True)
            headers = _headers(next(iterator, []))
            records = []
            for vals in iterator:
                if not any(v is not None for v in vals):
                    continue
                records.append({name: vals[idx] if idx < len(vals) else None for idx, name in enumerate(headers)})
                if len(records) >= limit:
                    break
            results.append((sheet_name, headers, records))
        return results
    finally:
        book.close()


def _sql_engine(kind, config):
    try:
        import sqlalchemy as sa
    except ImportError:
        raise DataError("DRIVER_MISSING", "SQL 데이터원에는 SQLAlchemy 및 해당 DB 드라이버가 필요합니다.") from None
    host = config.get("host")
    if not isinstance(host, str) or not host or any(c in host for c in "\r\n\x00"):
        raise DataError("INVALID_CONFIG", "데이터베이스 서버 주소가 필요합니다.")
    username, password, database = config.get("username"), config.get("password"), config.get("database")
    port = _integer(config.get("port"), {"postgresql": 5432, "mariadb": 3306, "oracle": 1521, "mssql": 1433}[kind], 1, 65535)
    query, connect_args = {}, {}
    if kind == "postgresql":
        driver = "postgresql+psycopg"
        connect_args = {"connect_timeout": 10, "options": "-c statement_timeout=15000 -c default_transaction_read_only=on"}
        query["sslmode"] = config.get("sslmode", "require")
    elif kind == "mariadb":
        driver = "mysql+pymysql"
        connect_args = {"connect_timeout": 10, "read_timeout": 15, "write_timeout": 15, "charset": "utf8mb4"}
        if config.get("ssl_ca"):
            connect_args["ssl"] = {"ca": config["ssl_ca"], "check_hostname": True}
        elif config.get("tls", True):
            connect_args["ssl"] = ssl.create_default_context()
    elif kind == "oracle":
        driver = "oracle+oracledb"
        query["service_name"] = config.get("service_name") or database or ""
        database = None
        if config.get("mode", "thin") != "thin":
            raise DataError("UNSUPPORTED_CAPABILITY", "이 worker는 Oracle Thin 모드만 지원합니다. Thick 전용 worker를 별도 구성하세요.")
        connect_args["tcp_connect_timeout"] = 10
    else:
        driver = "mssql+pyodbc"
        query = {"driver": config.get("driver", "ODBC Driver 18 for SQL Server"), "Encrypt": "yes", "TrustServerCertificate": "no"}
        connect_args = {"timeout": 10}
    url = sa.URL.create(driver, username=username, password=password, host=host, port=port, database=database, query=query)
    engine = sa.create_engine(url, connect_args=connect_args, pool_pre_ping=True, pool_size=1, max_overflow=0, pool_timeout=10)
    return engine


def _sql_schema(kind, config):
    import sqlalchemy as sa
    engine = _sql_engine(kind, config)
    try:
        inspector = sa.inspect(engine)
        schema = config.get("schema") or inspector.default_schema_name
        objects = []
        for typ, names in (("table", inspector.get_table_names(schema=schema)), ("view", inspector.get_view_names(schema=schema))):
            for name in sorted(names):
                primary = set((inspector.get_pk_constraint(name, schema=schema) or {}).get("constrained_columns") or []) if typ == "table" else set()
                columns = [{"column": col["name"], "name": col["name"], "type": "datetime" if kind == "oracle" and str(col["type"]).upper() == "DATE" else _sql_type(col["type"]),
                            "native_type": str(col["type"]), "nullable": bool(col.get("nullable", True)),
                            "primary_key": col["name"] in primary, "comment": col.get("comment"), "inferred": False}
                           for col in inspector.get_columns(name, schema=schema)]
                objects.append(_object(name, columns, schema, typ))
        return objects
    finally:
        engine.dispose()


class _PinnedHTTPS(http.client.HTTPSConnection):
    def __init__(self, host, address, port=443):
        super().__init__(host, port=port, timeout=15, context=ssl.create_default_context())
        self._address = address

    def connect(self):
        raw = socket.create_connection((self._address, self.port), timeout=self.timeout)
        self.sock = self._context.wrap_socket(raw, server_hostname=self.host)


def _https_json(endpoint, headers=None, allowed_hosts=None, *, csv_text=False):
    """Exact host allowlist, public DNS/IP, TLS validation, no redirects or proxies.

    DNS is resolved once and the HTTPS socket is pinned to that validated address,
    preventing a second DNS lookup from rebinding the request to an internal host.
    allowed_hosts must be injected by server policy, never project definitions.
    """
    if not isinstance(endpoint, str) or any(c in endpoint for c in "\r\n\x00"):
        raise DataError("INVALID_CONFIG", "올바른 HTTPS 자료 주소가 필요합니다.")
    parsed = urlsplit(endpoint)
    host = (parsed.hostname or "").lower().rstrip(".")
    allowed = {str(h).lower().rstrip(".") for h in (allowed_hosts or [])}
    try:
        port = parsed.port or 443
    except ValueError:
        raise DataError("UNSAFE_ENDPOINT", "허용되지 않은 자료 주소입니다.") from None
    if parsed.scheme != "https" or not host or host not in allowed or parsed.username or parsed.password or parsed.fragment or port != 443:
        raise DataError("UNSAFE_ENDPOINT", "관리자가 허용한 고정 HTTPS 자료 주소만 사용할 수 있습니다.")
    addresses = {entry[4][0] for entry in socket.getaddrinfo(host, port, type=socket.SOCK_STREAM)}
    if not addresses or any(not ipaddress.ip_address(address).is_global for address in addresses):
        raise DataError("UNSAFE_ENDPOINT", "사설·예약 네트워크 자료 주소는 사용할 수 없습니다.")
    clean_headers = {"Accept": "application/json"}
    for key, value in (headers or {}).items():
        if not isinstance(key, str) or not isinstance(value, str) or any(x in key + value for x in "\r\n\x00") or key.lower() in {"host", "connection", "content-length", "transfer-encoding"}:
            raise DataError("INVALID_CONFIG", "REST 요청 헤더가 올바르지 않습니다.")
        clean_headers[key] = value
    conn = _PinnedHTTPS(host, sorted(addresses)[0], port)
    try:
        target = parsed.path or "/"
        if parsed.query:
            target += "?" + parsed.query
        conn.request("GET", target, headers=clean_headers)
        response = conn.getresponse()
        if response.status == 429:
            raise DataError("RATE_LIMIT", "자료 API 요청 한도에 도달했습니다. 잠시 후 다시 실행하세요.")
        if response.status != 200 and csv_text:
            raise DataError('AUTH_REQUIRED', '공개 시트를 읽을 수 없습니다. 링크 공개/다운로드 권한을 확인하거나 OAuth로 연결하세요.')
        if response.status != 200:
            raise DataError("REMOTE_ERROR", "자료 API가 정상 응답하지 않았습니다. 연결 권한과 고정 주소를 확인하세요.")
        raw = response.read(_SOURCE_BYTES + 1)
        if len(raw) > _SOURCE_BYTES:
            raise DataError("DATA_LIMIT", "자료 API 응답이 허용 크기를 초과했습니다.")
        if csv_text:
            content_type = response.getheader('Content-Type', '').split(';')[0].lower()
            if content_type not in {'text/csv', 'text/plain', 'application/csv'} or raw.lstrip().startswith(b'<'):
                raise DataError('AUTH_REQUIRED', '공개 CSV를 읽을 수 없습니다. 링크 공개 설정을 확인하거나 OAuth로 연결하세요.')
            return raw.decode('utf-8-sig')
        return json.loads(raw)
    finally:
        conn.close()


def _sheet_headers(config):
    token = config.get("access_token")
    if not token and config.get("service_account"):
        try:
            from google.oauth2.service_account import Credentials
            from google.auth.transport.requests import Request
        except ImportError:
            raise DataError("DRIVER_MISSING", "서비스 계정 인증에는 google-auth와 requests가 필요합니다.") from None
        account = config["service_account"]
        if isinstance(account, str):
            account = json.loads(account)
        if not isinstance(account, dict) or account.get("token_uri", "https://oauth2.googleapis.com/token") != "https://oauth2.googleapis.com/token":
            raise DataError("INVALID_CONFIG", "Google 서비스 계정의 토큰 주소가 올바르지 않습니다.")
        credentials = Credentials.from_service_account_info(account, scopes=["https://www.googleapis.com/auth/spreadsheets.readonly"])
        credentials.refresh(Request())
        token = credentials.token
    if not isinstance(token, str) or not token:
        raise DataError("AUTH_REQUIRED", "Google OAuth 액세스 토큰 또는 공유 권한이 있는 서비스 계정이 필요합니다.")
    return {"Authorization": "Bearer " + token}


def _sheet_base(config):
    spreadsheet_id = config.get("spreadsheet_id")
    if not isinstance(spreadsheet_id, str) or not re.fullmatch(r"[A-Za-z0-9_-]{8,200}", spreadsheet_id):
        raise DataError("INVALID_CONFIG", "Google 스프레드시트 식별자가 올바르지 않습니다.")
    return "https://sheets.googleapis.com/v4/spreadsheets/" + spreadsheet_id


def _public_sheet_numeric_hints(config):
    """Read declared Google column types; retain CSV strings for existing bindings."""
    query = {'tqx': 'out:json', 'headers': '1', 'tq': 'limit 0'}
    if config.get('sheet'):
        query['sheet'] = str(config['sheet'])
    else:
        query['gid'] = str(config.get('gid', '0'))
    endpoint = 'https://docs.google.com/spreadsheets/d/' + config['spreadsheet_id'] + '/gviz/tq?' + urlencode(query)
    text = _https_json(endpoint, {}, ['docs.google.com'], csv_text=True)
    match = re.fullmatch(r'\s*(?:/\*.*?\*/\s*)?google\.visualization\.Query\.setResponse\((\{.*\})\);?\s*', text, re.S)
    try:
        payload = json.loads(match.group(1) if match else text)
        if payload.get('status') not in {'ok', 'warning'}:
            raise ValueError
        columns = payload['table']['cols']
        if not isinstance(columns, list) or any(not isinstance(col, dict) for col in columns):
            raise ValueError
        return columns
    except (ValueError, KeyError, TypeError, AttributeError):
        raise DataError('INVALID_RESPONSE', '공개 시트의 컬럼 타입 정보를 확인할 수 없습니다.') from None


def _sheet_records(config, name, limit):
    header = _integer(config.get("header_row"), 1, 1, 10000)
    if config.get('auth_mode') == 'public':
        _sheet_base(config)
        if header != 1:
            raise DataError('INVALID_CONFIG', '공개 시트는 첫 행을 헤더로 사용합니다. 다른 헤더 행은 OAuth로 연결하세요.')
        query = {'tqx': 'out:csv', 'headers': '1'}
        if config.get('sheet'):
            query['sheet'] = str(config['sheet'])
        else:
            gid = str(config.get('gid', '0'))
            if not re.fullmatch(r'\d{1,20}', gid):
                raise DataError('INVALID_CONFIG', '시트 gid가 올바르지 않습니다.')
            query['gid'] = gid
        endpoint = 'https://docs.google.com/spreadsheets/d/' + config['spreadsheet_id'] + '/gviz/tq?' + urlencode(query)
        raw = _https_json(endpoint, {}, ['docs.google.com'], csv_text=True)
        values = list(csv.reader(io.StringIO(raw)))[header - 1:]
        return _sheet_values(values, limit)
    # Escaping quotes in A1 avoids interpreting a sheet name as a range formula.
    # Read the complete bounded API response, not a row window: blank rows must
    # not cause a truncated source to look like a complete aggregate snapshot.
    span = "'" + name.replace("'", "''") + "'"
    payload = _https_json(_sheet_base(config) + "/values/" + quote(span, safe="") + "?valueRenderOption=UNFORMATTED_VALUE&dateTimeRenderOption=FORMATTED_STRING", _sheet_headers(config), ["sheets.googleapis.com"])
    values = payload.get("values", [])[header - 1:]
    return _sheet_values(values, limit)


def _sheet_values(values, limit):
    names = _headers(values[0] if values else [])
    if len(names) > 4096:
        raise DataError("DATA_LIMIT", "Google Sheets 열 수가 허용 범위를 초과했습니다.")
    rows = []
    for vals in values[1:]:
        if not any(v is not None and v != "" for v in vals):
            continue
        if len(vals) > len(names):
            raise DataError("INVALID_FILE", "Google Sheets 데이터 열 수가 헤더보다 많습니다.")
        rows.append({key: vals[idx] if idx < len(vals) and vals[idx] != "" else None for idx, key in enumerate(names)})
        if len(rows) >= limit:
            break
    return names, rows


def _rest_records(config, limit):
    payload = _https_json(config.get("endpoint"), config.get("headers"), config.get("allowed_hosts"))
    path = config.get("records_path", "")
    if path:
        if not isinstance(path, str) or not re.fullmatch(r"[A-Za-z0-9_-]+(?:\.[A-Za-z0-9_-]+)*", path):
            raise DataError("INVALID_CONFIG", "JSON 목록 경로가 올바르지 않습니다.")
        for part in path.split("."):
            if not isinstance(payload, dict) or part not in payload:
                raise DataError("INVALID_FILE", "REST 응답에서 설정된 자료 목록을 찾을 수 없습니다.")
            payload = payload[part]
    if not isinstance(payload, list) or any(not isinstance(row, dict) for row in payload):
        raise DataError("INVALID_FILE", "REST 자료는 JSON 객체의 배열이어야 합니다.")
    names = list(dict.fromkeys(key for row in payload[:100] for key in row))
    if any(not isinstance(name, str) or not name for name in names):
        raise DataError("INVALID_FILE", "REST 자료의 필드 이름이 올바르지 않습니다.")
    # Never silently omit a field appearing after the schema sample.
    if any(set(row) - set(names) for row in payload):
        raise DataError("SCHEMA_CHANGED", "REST 자료의 필드 구성이 표본 이후 달라졌습니다.")
    return names, [{name: row.get(name) for name in names} for row in payload[:limit]]


def _mdb(command, path, *args):
    import tempfile
    executable = shutil.which(command)
    if executable is None:
        raise DataError("DRIVER_MISSING", "Access 파일을 읽으려면 서버에 MDBTools를 설치하세요.")
    # Stream subprocess output to a temporary file instead of unbounded RAM.
    with tempfile.TemporaryFile() as output:
        result = subprocess.run([executable, str(path), *args], stdout=output, stderr=subprocess.DEVNULL, timeout=15, check=False)
        if result.returncode != 0:
            raise DataError("UNSUPPORTED_FILE", "MDBTools가 이 Access 파일을 읽지 못했습니다. 암호화·파일 형식·드라이버 지원을 확인하세요.")
        output.seek(0)
        raw = output.read(_SOURCE_BYTES + 1)
        if len(raw) > _SOURCE_BYTES:
            raise DataError("DATA_LIMIT", "Access 추출 자료가 허용 크기를 초과했습니다.")
        return raw.decode("utf-8")


def _access_records(config, name, limit):
    import io
    path = _file(config, {".mdb", ".accdb"})
    # subprocess receives an argument array; a table name can never be a shell command.
    stream = csv.reader(io.StringIO(_mdb("mdb-export", path, name)))
    names = _headers(next(stream, []))
    rows = []
    for vals in stream:
        rows.append({key: vals[idx] if idx < len(vals) and vals[idx] != "" else None for idx, key in enumerate(names)})
        if len(rows) >= limit:
            break
    return names, rows


def introspect(kind, config):
    """Return physical objects/columns; sample inference never becomes a DB fact."""
    kind = _kind(kind)
    try:
        if kind == "sqlite":
            objects = _sqlite_schema(config)
        elif kind == "csv":
            names, rows = _csv_records(config, 100)
            objects = [_object("data", _columns(names, rows), typ="file")]
        elif kind == "excel":
            objects = [_object(name, _columns(names, rows), typ="sheet", warning="formula_cached_values_only")
                       for name, names, rows in _excel_records(config, 100)]
        elif kind in {"postgresql", "mariadb", "oracle", "mssql"}:
            objects = _sql_schema(kind, config)
        elif kind == "google_sheets":
            if config.get('auth_mode') == 'public':
                name = config.get('sheet') or '공개 시트'
                names, rows = _sheet_records(config, name, 100)
                columns = _columns(names, rows)
                try:
                    hints = _public_sheet_numeric_hints(config)
                except DataError:
                    # Type hints are optional; a readable CSV remains usable as before.
                    hints = []
                if len(hints) == len(columns):
                    for column, hint in zip(columns, hints):
                        if hint.get('type') == 'number' and isinstance(hint.get('label'), str) and hint['label'].strip() == column['label']:
                            column.update(source_type='number', suggested_type='decimal', suggested_conversion='to_decimal')
                return {'objects': [_object(name, columns, typ='sheet')]}
            payload = _https_json(_sheet_base(config) + "?fields=sheets.properties", _sheet_headers(config), ["sheets.googleapis.com"])
            objects = []
            for sheet in payload.get("sheets", []):
                name = sheet["properties"]["title"]
                names, rows = _sheet_records(config, name, 100)
                objects.append(_object(name, _columns(names, rows), typ="sheet"))
        elif kind == "rest":
            names, rows = _rest_records(config, 100)
            objects = [_object("data", _columns(names, rows), typ="api")]
        elif kind == "access":
            path = _file(config, {".mdb", ".accdb"})
            objects = []
            for name in _mdb("mdb-tables", path, "-1").splitlines():
                if not name:
                    continue
                names, rows = _access_records(config, name, 100)
                objects.append(_object(name, _columns(names, rows), typ="table", warning="mdbtools_support_depends_on_file"))
        else:
            return _PLUGIN_CONNECTORS[kind].introspect(config)
        return {"objects": objects, "capabilities": dict(_BASE_CAPABILITIES),
                "verified": False, "verification": "external_connection_not_verified", "sample_limit": 100}
    except DataError:
        raise
    except ImportError:
        raise DataError("DRIVER_MISSING", "선택한 자료 연결에 필요한 드라이버를 설치하세요.") from None
    except Exception:
        raise DataError("CONNECTION_FAILED", "자료를 읽지 못했습니다. 연결 설정, 권한, 파일 형식을 확인하세요.") from None


def test_connection(kind, config):
    schema = introspect(kind, config)
    return {"connected": True, "ok": True, "object_count": len(schema["objects"]),
            "message": "읽기 연결을 확인했습니다.", "capabilities": schema.get("capabilities", dict(_BASE_CAPABILITIES))}


def _read(kind, config, mapping, columns, limit):
    name, schema = mapping["object"], mapping.get("schema")
    if kind == "sqlite":
        with _sqlite(config) as conn:
            query = "SELECT " + ",".join(_identifier(c) for c in columns) + " FROM " + _identifier(name) + " LIMIT ?"
            types = {col["name"]: _sql_type(col["type"]) for col in conn.execute("PRAGMA table_info(" + _identifier(name) + ")")}
            rows = [dict(row) for row in conn.execute(query, (limit,))]
            for row in rows:
                for column, value in row.items():
                    if value is not None and isinstance(value, str) and types.get(column) in {"date", "datetime"}:
                        try:
                            row[column] = dt.date.fromisoformat(value) if types[column] == "date" else dt.datetime.fromisoformat(value)
                        except ValueError:
                            raise DataError("TYPE_MISMATCH", "SQLite 날짜 컬럼에 올바르지 않은 날짜 값이 있습니다.") from None
            return rows
    if kind == "csv":
        return _csv_records(config, limit)[1]
    if kind == "excel":
        return _excel_records(config, limit, name)[0][2]
    if kind == "google_sheets":
        return _sheet_records(config, name, limit)[1]
    if kind == "rest":
        return _rest_records(config, limit)[1]
    if kind == "access":
        return _access_records(config, name, limit)[1]
    if kind in _PLUGIN_CONNECTORS:
        return _PLUGIN_CONNECTORS[kind].read(config, mapping, columns, limit)
    import sqlalchemy as sa
    engine = _sql_engine(kind, config)
    try:
        with engine.connect() as conn:
            if kind == "postgresql":
                conn.exec_driver_sql("SET TRANSACTION READ ONLY")
            elif kind == "mariadb":
                conn.exec_driver_sql("SET TRANSACTION READ ONLY")
            elif kind == "oracle":
                conn.exec_driver_sql("SET TRANSACTION READ ONLY")
            if kind == "mssql":
                conn.connection.driver_connection.timeout = 15
            if kind == "oracle":
                conn.connection.driver_connection.call_timeout = 15000
            table = sa.Table(name, sa.MetaData(), *(sa.Column(c) for c in columns), schema=schema)
            result = conn.execution_options(stream_results=True).execute(sa.select(*table.c).limit(limit))
            return [dict(row) for row in result.mappings()]
    finally:
        engine.dispose()


def _convert(value, conversion, field):
    if value is None:
        if field.get("nullable", True) is False:
            raise DataError("NULL_VALUE", "NULL을 허용하지 않는 논리 필드에 빈 값이 있습니다.")
        return None
    try:
        if conversion in {None, "identity"}:
            return value
        if conversion == "to_string":
            return str(value)
        if conversion == "to_decimal":
            result = Decimal(str(value))
            if not result.is_finite():
                raise ValueError
            return result
        if conversion == "to_integer":
            result = Decimal(str(value))
            if not result.is_finite() or result != result.to_integral_value():
                raise ValueError
            return int(result)
        if conversion == "to_float":
            result = float(value)
            if not math.isfinite(result):
                raise ValueError
            return result
        if conversion == "to_boolean":
            if value in (True, 1, "1", "true", "True"):
                return True
            if value in (False, 0, "0", "false", "False"):
                return False
            raise ValueError
        if conversion == "to_date":
            return dt.date.fromisoformat(str(value))
        if conversion == "to_datetime":
            return dt.datetime.fromisoformat(str(value))
    except (ValueError, TypeError, InvalidOperation, OverflowError):
        raise DataError("CONVERSION_FAILED", "자료 값이 지정한 논리 필드 변환 규칙에 맞지 않습니다.") from None
    raise DataError("INVALID_MAPPING", "지원하지 않는 필드 변환 규칙입니다.")


def _coerce_filter(value, actual):
    if actual is None or value is None:
        return value
    try:
        if isinstance(actual, bool):
            return _convert(value, "to_boolean", {})
        if isinstance(actual, Decimal):
            return _convert(value, "to_decimal", {})
        if isinstance(actual, int):
            return _convert(value, "to_integer", {})
        if isinstance(actual, float):
            return _convert(value, "to_float", {})
        if isinstance(actual, dt.datetime):
            return value if isinstance(value, dt.datetime) else dt.datetime.fromisoformat(str(value))
        if isinstance(actual, dt.date):
            return value if isinstance(value, dt.date) else dt.date.fromisoformat(str(value))
        if isinstance(actual, str):
            return str(value)
        return value
    except (ValueError, InvalidOperation, TypeError):
        raise DataError("INVALID_FILTER", "필터 값의 자료형이 올바르지 않습니다.") from None


def _filter(node, fields, parameters, depth=0, budget=None):
    budget = [0] if budget is None else budget
    budget[0] += 1
    if budget[0] > 1000:
        raise DataError("INVALID_QUERY", "필터 노드 수가 허용 범위를 초과했습니다.")
    if not isinstance(node, dict) or depth > 10:
        raise DataError("INVALID_QUERY", "필터 구조가 올바르지 않거나 너무 깊습니다.")
    if "op" in node:
        op, items = node.get("op"), node.get("items", [])
        if op not in {"and", "or"} or not isinstance(items, list) or len(items) > 100:
            raise DataError("INVALID_QUERY", "필터 그룹이 올바르지 않습니다.")
        children = [_filter(item, fields, parameters, depth + 1, budget) for item in items]
        return (lambda row: all(child(row) for child in children)) if op == "and" else (lambda row: any(child(row) for child in children))
    field_id, op = node.get("field_id"), node.get("operator")
    valid_ops = {"eq", "ne", "gt", "gte", "lt", "lte", "contains", "starts_with", "ends_with", "in", "not_in", "is_null", "is_not_null", "between"}
    if field_id not in fields or op not in valid_ops:
        raise DataError("INVALID_QUERY", "허용되지 않은 필드 또는 필터 연산입니다.")
    if "parameter" in node:
        name = node["parameter"]
        if not isinstance(name, str) or name not in parameters:
            raise DataError("MISSING_PARAMETER", "필터에 필요한 매개변수가 없습니다.")
        value = parameters[name]
    else:
        value = node.get("value")
    if op in {"in", "not_in", "between"} and (not isinstance(value, list) or len(value) > 1000 or op == "between" and len(value) != 2):
        raise DataError("INVALID_FILTER", "목록 또는 범위 필터 값이 올바르지 않습니다.")
    comparisons = {"eq": operator.eq, "ne": operator.ne, "gt": operator.gt, "gte": operator.ge, "lt": operator.lt, "lte": operator.le}
    def predicate(row):
        actual = row.get(field_id)
        if op == "is_null":
            return actual is None
        if op == "is_not_null":
            return actual is not None
        if actual is None or value is None:
            return False  # SQL-style NULL comparison; use explicit null operators.
        try:
            if op in {"in", "not_in"}:
                found = actual in [_coerce_filter(v, actual) for v in value]
                return found if op == "in" else not found
            if op == "between":
                return _coerce_filter(value[0], actual) <= actual <= _coerce_filter(value[1], actual)
            expected = _coerce_filter(value, actual)
            if op in comparisons:
                return comparisons[op](actual, expected)
            if not isinstance(actual, str):
                raise DataError("INVALID_FILTER", "문자열 필터에는 문자열 필드를 사용하세요.")
            return str(expected) in actual if op == "contains" else actual.startswith(str(expected)) if op == "starts_with" else actual.endswith(str(expected))
        except TypeError:
            raise DataError("INVALID_FILTER", "비교할 수 없는 자료형의 필터입니다.") from None
    return predicate


def _sort(rows, sorts, valid_fields):
    if not isinstance(sorts, list) or len(sorts) > 20:
        raise DataError("INVALID_QUERY", "정렬 설정이 올바르지 않습니다.")
    for spec in reversed(sorts):
        if not isinstance(spec, dict) or spec.get("field_id") not in valid_fields or spec.get("direction", "asc") not in {"asc", "desc"} or spec.get("nulls", "last") not in {"first", "last"}:
            raise DataError("INVALID_QUERY", "정렬 필드 또는 방향이 올바르지 않습니다.")
        field, descending, null_first = spec["field_id"], spec.get("direction") == "desc", spec.get("nulls") == "first"
        def compare(a, b):
            x, y = a.get(field), b.get(field)
            if x is None or y is None:
                return 0 if x is None and y is None else (-1 if x is None else 1) * (1 if null_first else -1)
            try:
                result = (x > y) - (x < y)
            except TypeError:
                raise DataError("INVALID_QUERY", "정렬 필드의 자료형이 일관되지 않습니다.") from None
            return -result if descending else result
        rows.sort(key=functools.cmp_to_key(compare))
    return rows


def _aggregate(rows, groups, aggregates, fields):
    if not isinstance(groups, list) or not isinstance(aggregates, list) or len(groups) > 20 or len(aggregates) > 50:
        raise DataError("INVALID_QUERY", "그룹 또는 집계 설정이 올바르지 않습니다.")
    group_ids = [g.get("field_id") if isinstance(g, dict) else g for g in groups]
    if any(g not in fields for g in group_ids) or len(set(group_ids)) != len(group_ids):
        raise DataError("INVALID_QUERY", "그룹 필드가 올바르지 않습니다.")
    aggregate_ids, output_fields = set(), [fields[g] for g in group_ids]
    for spec in aggregates:
        if not isinstance(spec, dict):
            raise DataError("INVALID_QUERY", "집계 설정이 올바르지 않습니다.")
        aid, function, field = spec.get("aggregate_id"), spec.get("function"), spec.get("field_id")
        if not isinstance(aid, str) or not aid or aid in fields or aid in aggregate_ids or function not in {"sum", "count", "avg", "min", "max", "count_distinct"} or (field not in fields and not (function == "count" and field is None)):
            raise DataError("INVALID_QUERY", "집계 식별자, 함수 또는 필드가 올바르지 않습니다.")
        aggregate_ids.add(aid)
        output_fields.append({"field_id": aid, "alias": aid, "label": spec.get("label", aid), "type": "integer" if function in {"count", "count_distinct"} else "decimal" if function in {"sum", "avg"} else fields[field].get("type", "unknown"), "nullable": function not in {"count", "count_distinct"}})
    buckets = {}
    for row in rows:
        key = tuple(row.get(g) for g in group_ids)
        try:
            buckets.setdefault(key, []).append(row)
        except TypeError:
            raise DataError("INVALID_QUERY", "JSON·바이너리 필드는 그룹 키로 사용할 수 없습니다.") from None
    if not group_ids and not rows:
        buckets[()] = []
    result = []
    for key, members in buckets.items():
        output = dict(zip(group_ids, key))
        for spec in aggregates:
            field, fn = spec.get("field_id"), spec["function"]
            values = [row.get(field) for row in members if row.get(field) is not None] if field else []
            try:
                if fn == "count":
                    value = len(values) if field else len(members)
                elif fn == "count_distinct":
                    value = len(set(values))
                elif not values:
                    value = None
                elif fn in {"sum", "avg"}:
                    if any(isinstance(v, bool) or not isinstance(v, (int, float, Decimal)) for v in values):
                        raise DataError("INVALID_QUERY", "합계·평균에는 명시적으로 숫자 자료형으로 변환한 필드가 필요합니다.")
                    numbers = [Decimal(str(v)) for v in values]
                    value = sum(numbers) if fn == "sum" else sum(numbers) / len(numbers)
                else:
                    value = min(values) if fn == "min" else max(values)
            except (TypeError, InvalidOperation):
                raise DataError("INVALID_QUERY", "집계 필드 값의 자료형이 일관되지 않습니다.") from None
            output[spec["aggregate_id"]] = value
        result.append(output)
    return result, output_fields


def _json_safe(value):
    if value is None or isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            raise DataError("INVALID_VALUE", "자료에 유한하지 않은 숫자가 있습니다.")
        return value
    if isinstance(value, Decimal):
        if not value.is_finite():
            raise DataError("INVALID_VALUE", "자료에 유한하지 않은 숫자가 있습니다.")
        return str(value)
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, (bytes, bytearray, memoryview)):
        return base64.b64encode(bytes(value)).decode("ascii")
    if isinstance(value, dict):
        return {str(k): _json_safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(v) for v in value]
    return str(value)


def _check_logical(value, field):
    if value is None:
        return
    logical = field.get("type", "unknown")
    numeric = isinstance(value, (int, float, Decimal)) and not isinstance(value, bool)
    valid = {
        "string": isinstance(value, str),
        "integer": isinstance(value, int) and not isinstance(value, bool),
        "decimal": numeric,
        "float": numeric,
        "number": numeric,
        "boolean": isinstance(value, bool),
        "date": isinstance(value, dt.date) and not isinstance(value, dt.datetime),
        "datetime": isinstance(value, dt.datetime),
        "binary": isinstance(value, (bytes, bytearray, memoryview)),
        "json": isinstance(value, (dict, list)),
        "unknown": True,
    }.get(logical, False)
    if not valid:
        raise DataError("TYPE_MISMATCH", "자료 값이 논리 필드 자료형에 맞지 않습니다. 명시적 변환을 설정하세요.")


def execute_dataset(kind, config, contract, binding, parameters=None, max_rows=10000):
    """One bounded source snapshot -> mapping -> filter -> group -> sort -> projection.

    If a source exceeds the cap, fail instead of returning a misleading partial
    total. v0.1 deliberately disallows joins/raw SQL and query pushdown.
    """
    kind = _kind(kind)
    max_rows = _integer(max_rows, 10000, 1, _MAX_QUERY_ROWS)
    parameters = parameters or {}
    try:
        if not isinstance(contract, dict) or not isinstance(binding, dict):
            raise DataError("INVALID_MAPPING", "논리 데이터 계약과 자료 매핑이 필요합니다.")
        objects, fields, query = contract.get("objects", []), contract.get("fields", []), contract.get("query", {})
        if len(objects) != 1 or query.get("joins"):
            raise DataError("UNSUPPORTED_CAPABILITY", "현재 버전은 단일 테이블만 조회합니다. JOIN은 지원하지 않습니다.")
        if query.get("sql") or query.get("raw_sql"):
            raise DataError("UNSUPPORTED_CAPABILITY", "임의 SQL은 지원하지 않습니다.")
        if binding.get("dataset_id") != contract.get("dataset_id"):
            raise DataError("INVALID_MAPPING", "매핑의 논리 데이터셋 식별자가 일치하지 않습니다.")
        object_id = objects[0]["object_id"]
        mapping = binding.get("object_mappings", {}).get(object_id)
        if not isinstance(mapping, dict) or not isinstance(mapping.get("object"), str):
            raise DataError("INVALID_MAPPING", "물리 테이블 또는 시트 매핑이 필요합니다.")
        schema = introspect(kind, config)
        physical = next((obj for obj in schema["objects"] if obj["object"] == mapping["object"] and (obj.get("schema") or None) == (mapping.get("schema") or None)), None)
        if physical is None:
            raise DataError("SCHEMA_CHANGED", "매핑된 테이블·시트가 없습니다. 자료 매핑을 다시 확인하세요.")
        physical_columns = {col["column"]: col for col in physical["columns"]}
        field_map, selected_columns = {}, []
        if not fields or len(fields) > 500:
            raise DataError("INVALID_QUERY", "논리 필드는 1개 이상 500개 이하로 지정하세요.")
        for field in fields:
            fid = field.get("field_id")
            if not isinstance(fid, str) or not fid or fid in field_map or field.get("object_id") != object_id:
                raise DataError("INVALID_MAPPING", "논리 필드 식별자 또는 소속 테이블이 올바르지 않습니다.")
            fmap = binding.get("field_mappings", {}).get(fid)
            if fmap is None and not field.get("required", True):
                field_map[fid] = (field, None)
                continue
            if not isinstance(fmap, dict) or fmap.get("column") not in physical_columns:
                raise DataError("SCHEMA_CHANGED", "필수 컬럼이 없거나 컬럼 매핑이 올바르지 않습니다.")
            if fmap.get("conversion", "identity") not in {"identity", "to_string", "to_integer", "to_decimal", "to_float", "to_boolean", "to_date", "to_datetime"}:
                raise DataError("INVALID_MAPPING", "지원하지 않는 필드 변환 규칙입니다.")
            native = physical_columns[fmap["column"]]["type"]
            logical = field.get("type", "unknown")
            compatible = native == logical or native == "unknown" or logical == "unknown" or (
                native in {"integer", "decimal", "float"} and logical == "number"
            ) or (native == "integer" and logical in {"decimal", "float"})
            if fmap.get("conversion", "identity") == "identity" and not compatible:
                raise DataError("TYPE_MISMATCH", "물리·논리 필드 자료형이 다릅니다. 명시적 변환을 설정하세요.")
            selected_columns.append(fmap["column"])
            field_map[fid] = (field, fmap)
        if not selected_columns:
            raise DataError("INVALID_MAPPING", "조회할 물리 필드를 하나 이상 연결하세요.")
        raw_rows = _read(kind, config, mapping, list(dict.fromkeys(selected_columns)), max_rows + 1)
        if len(raw_rows) > max_rows:
            raise DataError("DATA_LIMIT", "자료 원본이 조회 한도를 초과했습니다. 전용 뷰·시트·파일로 범위를 줄이세요. 부분 집계는 반환하지 않습니다.")
        rows = []
        for raw in raw_rows:
            row = {}
            for fid, (field, fmap) in field_map.items():
                if fmap and fmap["column"] not in raw:
                    raise DataError("SCHEMA_CHANGED", "읽는 동안 자료의 컬럼 구성이 바뀌었습니다. 매핑을 다시 확인하세요.")
                value = raw.get(fmap["column"]) if fmap else field.get("default")
                row[fid] = _convert(value, fmap.get("conversion", "identity") if fmap else "identity", field)
                _check_logical(row[fid], field)
            rows.append(row)
        field_defs = {fid: field for fid, (field, _) in field_map.items()}
        predicate = _filter(query.get("filters", {"op": "and", "items": []}), field_defs, parameters)
        rows = [row for row in rows if predicate(row)]
        groups, aggregates = query.get("groups", []), query.get("aggregates", [])
        if groups or aggregates:
            rows, output_fields = _aggregate(rows, groups, aggregates, field_defs)
            output_map = {field["field_id"]: field for field in output_fields}
        else:
            output_map = field_defs
        rows = _sort(rows, query.get("sorts", []), output_map)
        projection = query.get("projection") or list(output_map)
        if not isinstance(projection, list) or len(projection) != len(set(projection)) or any(fid not in output_map for fid in projection):
            raise DataError("INVALID_QUERY", "출력 필드 구성이 올바르지 않습니다.")
        rows = [{fid: _json_safe(row.get(fid)) for fid in projection} for row in rows]
        if contract.get("cardinality") == "one" and len(rows) > 1:
            raise DataError("CARDINALITY_MISMATCH", "단일 행 데이터셋에 여러 행이 반환되었습니다.")
        return {"dataset_id": contract.get("dataset_id"), "fields": [output_map[fid] for fid in projection],
                "rows": rows, "row_count": len(rows), "source_row_count": len(raw_rows), "truncated": False,
                "generated_at": dt.datetime.now(dt.timezone.utc).isoformat(),
                "capabilities": dict(_BASE_CAPABILITIES), "execution_mode": "bounded_snapshot",
                "external_connection_verified": False}
    except DataError:
        raise
    except ImportError:
        raise DataError("DRIVER_MISSING", "선택한 자료 연결에 필요한 드라이버를 설치하세요.") from None
    except Exception:
        raise DataError("QUERY_FAILED", "자료를 조회하지 못했습니다. 연결 권한, 논리 매핑, 조회 설정을 확인하세요.") from None
