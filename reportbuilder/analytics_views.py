"""Owner/staff session-only statistics and exports; no bearer-token scopes."""
import html
import io
from functools import wraps

from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404
from django.views.decorators.http import require_GET
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .analytics import AnalyticsError, PDF_DETAIL_LIMIT, XLSX_DETAIL_LIMIT, statistics
from .models import Report
from .rendering import pdf_bytes


def _error(code, message, status):
    response = JsonResponse({"code": code, "message": message}, status=status)
    response["Cache-Control"] = "private, no-store"
    return response


def owner_session(function):
    @wraps(function)
    def wrapped(request, report_id, *args, **kwargs):
        if request.headers.get("Authorization") or getattr(request, "api_token", None) is not None:
            return _error("SESSION_REQUIRED", "접근 통계는 소유자·관리자의 로그인 세션으로만 조회할 수 있습니다.", 403)
        if not request.user.is_authenticated or not request.user.is_active or not request.session.get("_auth_user_id"):
            return _error("AUTH_REQUIRED", "접근 통계를 보려면 로그인하세요.", 401)
        queryset = Report.objects.all() if request.user.is_staff else Report.objects.filter(owner=request.user)
        report = get_object_or_404(queryset, pk=report_id)
        try:
            response = function(request, report, *args, **kwargs)
        except AnalyticsError as exc:
            return _error(exc.code, str(exc), 400)
        response["Cache-Control"] = "private, no-store"
        return response
    return wrapped


@require_GET
@owner_session
def statistics_view(request, report):
    return JsonResponse(statistics(report, request.GET))


def _safe_cell(value):
    text = str(value) if value is not None else ""
    if text.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + text
    return value if isinstance(value, (int, float, bool)) else text


def xlsx_statistics(data):
    book = Workbook()
    book.remove(book.active)
    def sheet(name, headers, rows):
        tab = book.create_sheet(name)
        tab.append([_safe_cell(value) for value in headers])
        for row in rows:
            tab.append([_safe_cell(value) for value in row])
        tab.freeze_panes = "A2"
        tab.auto_filter.ref = tab.dimensions
        for cell in tab[1]:
            cell.fill = PatternFill("solid", fgColor="16324F")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center")
        for column in tab.columns:
            letter = column[0].column_letter
            width = min(60, max(16, max(len(str(cell.value or "")) for cell in column) + 3))
            tab.column_dimensions[letter].width = width
        return tab
    summary = [
        ("보고서", data["report"]["name"]), ("시작일", data["range"]["start"]), ("종료일", data["range"]["end"]),
        ("시간대", "Asia/Seoul"), ("전체 이벤트 수", data["total"]),
        ("집계 기준", "이벤트 수 (고유 방문자 수가 아님)"),
        ("상세 포함 건수", len(data["recent"])), ("상세 제외 건수", data["details_omitted"]),
        ("상세 최대 건수", data["recent_limit"]), ("집계 범위", "요청 기간·필터 전체, 상세만 최근 건수 한도 적용"),
        ("국가 조회", "오프라인 GeoIP 사용" if data["meta"]["geolocation"]["available"] else "GeoIP 미설치: 공인 IP 국가 Unknown"),
        ("Private 의미", "사설·루프백·예약 IP (원본 IP 미저장)"),
    ]
    summary.extend((f"이벤트 {row['event']}", row["count"]) for row in data["event_counts"])
    summary.extend((f"필터 {key}", value or "전체") for key, value in data["filters"].items())
    sheet("Summary", ["항목", "값"], summary)
    sheet("Daily", ["날짜 (KST)", "이벤트 수"], [(row["date"], row["count"]) for row in data["daily"]])
    sheet("Country", ["국가 / 구분", "이벤트 수"], [(row["country"], row["count"]) for row in data["countries"]])
    sheet("Device", ["기기", "이벤트 수"], [(row["device"], row["count"]) for row in data["devices"]])
    sheet("Browser", ["브라우저", "이벤트 수"], [(row["browser"], row["count"]) for row in data["browsers"]])
    sheet("Details", ["접근 시각 (KST)", "이벤트", "국가", "기기", "브라우저", "OS"],
          [[row[key] for key in ("timestamp", "event", "country", "device", "browser", "os")] for row in data["recent"]])
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def statistics_pdf_html(data):
    """Explicit A4 sections; conservative bounded table rows per page."""
    def escaped(value):
        return html.escape(str(value), quote=True)
    pages = []
    def section(title, headers, rows):
        rows = list(rows)
        chunks = [rows[index:index + 30] for index in range(0, len(rows), 30)] or [[]]
        for index, chunk in enumerate(chunks):
            heading = title + (f" ({index + 1}/{len(chunks)})" if len(chunks) > 1 else "")
            table = "<table><thead><tr>" + "".join(f"<th>{escaped(value)}</th>" for value in headers) + "</tr></thead><tbody>"
            table += "".join("<tr>" + "".join(f"<td>{escaped(value)}</td>" for value in row) + "</tr>" for row in chunk)
            table += "</tbody></table>" if chunk else '<tr><td colspan="' + str(len(headers)) + '">조회 결과 없음</td></tr></tbody></table>'
            pages.append((heading, table))
    summary = [
        ("기간 (KST)", f"{data['range']['start']} ~ {data['range']['end']}"),
        ("전체 이벤트 수", data["total"]), ("집계 기준", "이벤트 수이며 고유 방문자 수가 아닙니다."),
        ("상세 포함 / 제외", f"{len(data['recent'])} / {data['details_omitted']}"),
        ("상세 한도", f"최근 {data['recent_limit']}건, 집계는 기간 전체"),
        ("국가 조회", "오프라인 GeoIP 사용" if data["meta"]["geolocation"]["available"] else "국가 조회 데이터베이스 미설치"),
        ("국가 구분", "Unknown: 국가 미확인 / Private: 사설·예약 IP"),
        ("개인정보", "원본 IP·사용자 에이전트·사용자 이름을 저장하지 않습니다."),
    ]
    summary.extend((f"이벤트 {row['event']}", row["count"]) for row in data["event_counts"])
    summary.extend((f"필터 {key}", value or "전체") for key, value in data["filters"].items())
    section("접근 통계 요약", ["항목", "값"], summary)
    section("일별 이벤트", ["날짜 (KST)", "이벤트 수"], [(row["date"], row["count"]) for row in data["daily"]])
    section("국가별 이벤트", ["국가 / 구분", "이벤트 수"], [(row["country"], row["count"]) for row in data["countries"]])
    section("기기별 이벤트", ["기기", "이벤트 수"], [(row["device"], row["count"]) for row in data["devices"]])
    section("브라우저별 이벤트", ["브라우저", "이벤트 수"], [(row["browser"], row["count"]) for row in data["browsers"]])
    section("최근 접근 상세", ["접근 시각 (KST)", "이벤트", "국가", "기기", "브라우저", "OS"],
            [[row[key].replace("T", " ").replace("+09:00", "") if key == "timestamp" else row[key]
              for key in ("timestamp", "event", "country", "device", "browser", "os")] for row in data["recent"]])
    style = """@page{size:A4;margin:0}*{box-sizing:border-box}body{margin:0;color:#182638;font-family:'Noto Sans CJK KR','Noto Sans KR',Arial,sans-serif;font-size:9pt}.page{width:210mm;height:297mm;padding:13mm 14mm;position:relative;break-after:page;overflow:hidden}.page:last-child{break-after:auto}h1{font-size:15pt;margin:0 0 3mm;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}h2{font-size:12pt;margin:3mm 0}p{margin:0 0 4mm;color:#526477}table{width:100%;border-collapse:collapse;table-layout:fixed;font-size:8pt}th,td{border:0.2mm solid #d3dbe4;padding:1.3mm 1mm;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;height:6.6mm}th{background:#edf2f7;text-align:left}td:first-child,th:first-child{width:28%}.footer{position:absolute;bottom:8mm;left:14mm;right:14mm;color:#66778a;font-size:8pt;border-top:0.2mm solid #d3dbe4;padding-top:2mm;display:flex;justify-content:space-between}"""
    title = escaped(data["report"]["name"])
    subtitle = escaped(f"{data['range']['start']} ~ {data['range']['end']} · Asia/Seoul · 이벤트 수 {data['total']}")
    body = "".join(f'<section class="page"><h1>{title}</h1><p>{subtitle}</p><h2>{escaped(heading)}</h2>{table}<div class="footer"><span>원본 IP·UA 미저장 / 상세 제한 {data["recent_limit"]}건 · 제외 {data["details_omitted"]}건</span><span>{index + 1} / {len(pages)}</span></div></section>' for index, (heading, table) in enumerate(pages))
    return '<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta http-equiv="Content-Security-Policy" content="default-src \'none\'; style-src \'unsafe-inline\'; font-src \'none\'"><title>보고서 접근 통계</title><style>' + style + "</style></head><body>" + body + "</body></html>"


@require_GET
@owner_session
def statistics_export(request, report, format):
    if format not in {"xlsx", "pdf"}:
        return _error("UNSUPPORTED_FORMAT", "통계는 xlsx 또는 pdf로 내보낼 수 있습니다.", 400)
    data = statistics(report, request.GET, XLSX_DETAIL_LIMIT if format == "xlsx" else PDF_DETAIL_LIMIT)
    if format == "xlsx":
        content = xlsx_statistics(data)
        content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        try:
            content = pdf_bytes(statistics_pdf_html(data))
        except Exception:
            return _error("PDF_UNAVAILABLE", "PDF 생성기를 사용할 수 없습니다. 서버에 Playwright와 Chromium을 설치하고 다시 시도하세요.", 503)
        content_type = "application/pdf"
    response = HttpResponse(content, content_type=content_type)
    response["Content-Disposition"] = f'attachment; filename="report-statistics.{format}"'
    return response
