import copy
import csv
import io
import json
import logging
import uuid
from datetime import timedelta
from functools import wraps
from pathlib import Path

import markdown
from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core import signing
from django.core.management import call_command
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.base import ContentFile
from django.db import IntegrityError, transaction
from django.http import FileResponse, HttpResponse, Http404, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.csrf import csrf_exempt, ensure_csrf_cookie
from django.views.decorators.http import require_http_methods, require_POST
from django.views.decorators.clickjacking import xframe_options_exempt
from openpyxl import Workbook
from PIL import Image

from .analytics import record_report_access
from .data import DataError, connector_catalog, introspect, test_connection
from .definition import DefinitionError, default_definition, validate_definition
from .models import Asset, Connection, DemoSeed, EmbedNonce, Execution, Project, Publication, Report, Revision
from .packaging import MAX_ASSET, export_project, import_project, validate_asset
from .rendering import pdf_bytes
from .services import (audit, connections_for, editable, execute_report, policy_fingerprint,
                       reports_for, resolve_connections, run_datasets)

logger = logging.getLogger(__name__)


def api(function):
    @wraps(function)
    def wrapped(request, *args, **kwargs):
        if not request.user.is_authenticated:
            return JsonResponse({"code": "AUTH_REQUIRED", "message": "로그인이 필요합니다."}, status=401)
        try:
            token = getattr(request, "api_token", None)
            if token:
                scopes = {"report_api": "read" if request.method == "GET" else "edit", "preview_api": "edit",
                          "bindings_api": "edit", "publish_api": "edit", "execute_api": "run", "embed_session": "embed"}
                scope = scopes.get(function.__name__)
                report_id = kwargs.get("report_id")
                if function.__name__ == "embed_session":
                    report_id = body(request).get("report_id")
                if not scope or scope not in token.scopes or str(report_id) not in token.report_ids:
                    raise PermissionDenied("API 토큰의 보고서/기능 범위를 벗어났습니다.")
            return function(request, *args, **kwargs)
        except PermissionDenied as exc:
            return JsonResponse({"code": "ACCESS_DENIED", "message": str(exc)}, status=403)
        except (DataError, DefinitionError) as exc:
            return JsonResponse({"code": getattr(exc, "code", "VALIDATION_ERROR"), "message": str(exc)}, status=400)
        except Http404:
            return JsonResponse({"code": "NOT_FOUND", "message": "접근 가능한 항목을 찾을 수 없습니다."}, status=404)
        except (ValueError, KeyError, TypeError, ValidationError) as exc:
            logger.info("Invalid report request: %s", type(exc).__name__)
            return JsonResponse({"code": "INVALID_REQUEST", "message": "입력 형식 또는 참조를 확인해 주세요."}, status=400)
    return wrapped


def body(request):
    if len(request.body) > 5 * 1024 * 1024:
        raise ValueError("Definition too large")
    value = json.loads(request.body or b"{}")
    if not isinstance(value, dict):
        raise ValueError("Expected object")
    return value


def report_payload(report, user=None):
    publication = getattr(report, "publication", None)
    publication_url = reverse("publication", args=[publication.pk]) if publication and publication.enabled and report.enabled else None
    if user is not None and not (user.is_staff or report.owner_id == user.pk):
        # Viewer access authorizes a published version, never an in-progress draft
        # or local source mappings. Keep the same payload keys for API consumers.
        revision = publication.revision if publication and publication.enabled and report.enabled else None
        return {"id": str(report.pk), "name": report.name, "definition": revision.definition if revision else None,
                "revision": revision.number if revision else None, "bindings": [], "enabled": report.enabled,
                "published_revision": revision.number if revision else None, "publication_url": publication_url}
    return {"id": str(report.pk), "name": report.name, "definition": report.definition,
            "revision": report.revision, "bindings": report.bindings, "enabled": report.enabled,
            "published_revision": publication.revision.number if publication and publication.enabled else None,
            "publication_url": publication_url}


def new_revision(report):
    report.revision += 1
    report.save()
    Revision.objects.create(report=report, number=report.revision, definition=report.definition,
                            bindings=report.bindings)


@login_required
@ensure_csrf_cookie
def dashboard(request):
    reports = reports_for(request.user)
    return render(request, "reportbuilder/dashboard.html", {"reports": reports.order_by("-updated_at")[:6],
                  "report_count": reports.count(), "connection_count": connections_for(request.user).count(),
                  "executions": Execution.objects.filter(user=request.user).order_by("-created_at")[:8]})


@login_required
@ensure_csrf_cookie
def library(request):
    reports = reports_for(request.user).order_by("-updated_at")
    if request.GET.get("q"):
        reports = reports.filter(name__icontains=request.GET["q"])
    demo_pending = request.user.is_staff and not DemoSeed.objects.filter(owner=request.user, completed=True).exists() and not Report.objects.filter(owner=request.user, name="매출 현황 예제").exists()
    return render(request, "reportbuilder/library.html", {"reports": reports, "demo_pending": demo_pending})


@login_required
@require_POST
@transaction.atomic
def ensure_demo(request):
    if not request.user.is_staff or not request.user.is_active:
        raise PermissionDenied
    marker, _ = DemoSeed.objects.get_or_create(owner=request.user)
    marker = DemoSeed.objects.select_for_update().get(pk=marker.pk)
    if not marker.completed:
        call_command("seed_demo", username=request.user.username, stdout=io.StringIO())
    return JsonResponse({"ready": True})


@login_required
@ensure_csrf_cookie
def designer(request, report_id):
    report = get_object_or_404(reports_for(request.user), pk=report_id)
    editable(report, request.user)
    return render(request, "reportbuilder/designer.html", {"report": report,
                  "publication_url": report_payload(report, request.user)["publication_url"]})


@login_required
def viewer(request, report_id):
    report = get_object_or_404(reports_for(request.user), pk=report_id)
    response = render_viewer(request, report)
    if request.method == "GET":
        record_report_access(request, report, "view")
    return response


def render_viewer(request, report, published=False):
    can_edit = request.user.is_staff or report.owner_id == request.user.pk
    use_publication = published or not can_edit
    definition = report.definition
    if use_publication:
        publication = getattr(report, "publication", None)
        if not report.enabled or not publication or not publication.enabled:
            raise PermissionDenied("게시된 보고서가 없습니다.")
        definition = publication.revision.definition
    return render(request, "reportbuilder/viewer.html", {
        "report": report, "viewer_definition": definition,
        "published": use_publication, "can_edit": can_edit})


@api
@require_http_methods(["GET", "POST"])
def reports_api(request):
    if request.method == "GET":
        return JsonResponse({"reports": [report_payload(r, request.user) for r in reports_for(request.user)]})
    value = body(request)
    name = str(value.get("name", "새 보고서")).strip()[:200] or "새 보고서"
    definition = value.get("definition", default_definition(name))
    validate_definition(definition)
    with transaction.atomic():
        project = Project.objects.create(name=name, owner=request.user)
        report = Report.objects.create(name=name, owner=request.user, project=project, definition=definition)
        new_revision(report)
    audit(request.user, "report_create", report.pk)
    return JsonResponse(report_payload(report), status=201)


@api
@require_http_methods(["GET", "PUT", "DELETE"])
def report_api(request, report_id):
    report = get_object_or_404(reports_for(request.user), pk=report_id)
    if request.method == "GET":
        return JsonResponse(report_payload(report, request.user))
    editable(report, request.user)
    if request.method == "DELETE":
        if hasattr(report, "publication"):
            raise ValueError("Unpublish before deleting")
        report.delete()
        return JsonResponse({"deleted": True})
    value = body(request)
    definition = value.get("definition", report.definition)
    validate_definition(definition)
    with transaction.atomic():
        report = Report.objects.select_for_update().get(pk=report.pk)
        if value.get("expected_revision") != report.revision:
            return JsonResponse({"code": "REVISION_CONFLICT", "message": "다른 변경이 있습니다. 새로고침 후 다시 저장하세요."}, status=409)
        report.name = str(value.get("name", report.name))[:200]
        report.definition = definition
        new_revision(report)
    audit(request.user, "report_save", report.pk, revision=report.revision)
    return JsonResponse(report_payload(report))


def validate_bindings(user, report, bindings):
    if not isinstance(bindings, list) or len(bindings) > 20:
        raise ValueError("Invalid bindings")
    datasets = {d["dataset_id"]: d for d in report.definition.get("datasets", [])}
    if len({b["dataset_id"] for b in bindings}) != len(bindings):
        raise ValueError("Duplicate bindings")
    resolve_connections(user, bindings)
    for binding in bindings:
        if set(binding) - {"dataset_id", "connection_id", "object_mappings", "field_mappings"}:
            raise ValueError("Invalid binding keys")
        if binding["dataset_id"] not in datasets:
            raise ValueError("Unknown dataset")
        for m in binding.get("object_mappings", {}).values():
            if set(m) - {"schema", "object"}:
                raise ValueError("Invalid object mapping")
        for m in binding.get("field_mappings", {}).values():
            if set(m) - {"column", "conversion"}:
                raise ValueError("Invalid field mapping")


@api
@require_http_methods(["PUT"])
def bindings_api(request, report_id):
    report = get_object_or_404(reports_for(request.user), pk=report_id)
    editable(report, request.user)
    value = body(request)
    bindings = value.get("bindings", [])
    validate_bindings(request.user, report, bindings)
    with transaction.atomic():
        report = Report.objects.select_for_update().get(pk=report.pk)
        report.bindings = bindings
        new_revision(report)
    audit(request.user, "binding_save", report.pk)
    return JsonResponse(report_payload(report))


@api
@require_POST
def preview_api(request, report_id):
    report = get_object_or_404(reports_for(request.user), pk=report_id)
    execution, rendered = execute_report(report, request.user, body(request).get("parameters", {}))
    return JsonResponse({"execution_id": str(execution.pk), "html": rendered["html"],
                         "page_count": rendered["page_count"], "row_count": execution.row_count})


@api
@require_POST
def publish_api(request, report_id):
    report = get_object_or_404(reports_for(request.user), pk=report_id)
    editable(report, request.user)
    value = body(request)
    if value.get("enabled") is False:
        Publication.objects.filter(report=report).update(enabled=False)
        return JsonResponse({"enabled": False})
    number = int(value.get("revision", report.revision))
    revision = get_object_or_404(report.revisions, number=number)
    validate_definition(revision.definition)
    validate_bindings(request.user, report, revision.bindings)
    # Validate actual data bindings before exposing a revision. Parameters may be supplied for required filters.
    run_datasets(report, request.user, revision.definition, revision.bindings, value.get("parameters", {}))
    publication, _ = Publication.objects.update_or_create(report=report, defaults={"revision": revision, "enabled": True})
    audit(request.user, "publish", report.pk, revision=number)
    return JsonResponse({"publication_url": reverse("publication", args=[publication.pk]), "revision": number})


@api
@require_POST
def execute_api(request, report_id):
    report = get_object_or_404(reports_for(request.user), pk=report_id)
    execution, rendered = execute_report(report, request.user, body(request).get("parameters", {}), published=True)
    record_report_access(request, report, "execute")
    return JsonResponse({"execution_id": str(execution.pk), "html": rendered["html"],
                         "page_count": rendered["page_count"], "row_count": execution.row_count})


SECRET_FIELDS = {"password", "access_token", "refresh_token", "service_account", "headers", "api_key", "wallet_password"}
EDITABLE_SECRET_FIELDS = SECRET_FIELDS - {"access_token", "refresh_token"}
FILE_EXTENSIONS = {"excel": {".xlsx", ".xlsm"}, "csv": {".csv", ".tsv"},
                   "sqlite": {".sqlite", ".sqlite3", ".db"}, "access": {".mdb", ".accdb"}}


def create_connection(user, name, kind, config, upload=None):
    if not user.is_staff:
        raise PermissionDenied("데이터 연결은 관리자만 등록할 수 있습니다.")
    if not isinstance(config, dict) or len(json.dumps(config)) > 100000:
        raise ValueError("Invalid connection config")
    if kind not in {"excel", "csv", "sqlite", "access", "postgresql", "mariadb", "mysql", "oracle", "mssql", "sheets", "google_sheets", "rest"}:
        raise ValueError("Unknown connector")
    if set(config) & {"path", "allowed_hosts", "url", "connection_string"}:
        raise ValueError("Forbidden connection config")
    if kind in FILE_EXTENSIONS:
        if not upload or upload.size > 50 * 1024 * 1024 or Path(upload.name).suffix.lower() not in FILE_EXTENSIONS[kind]:
            raise ValueError("File type/size mismatch")
    secrets = {k: config[k] for k in SECRET_FIELDS if k in config}
    public_config = {k: v for k, v in config.items() if k not in SECRET_FIELDS}
    connection = Connection(name=str(name)[:200], kind=kind, config=public_config, owner=user)
    connection.set_secrets(secrets)
    if upload:
        connection.upload = upload
    connection.save()
    audit(user, "connection_create", connection.pk, connector=kind)
    return connection


@login_required
@ensure_csrf_cookie
@require_http_methods(["GET", "POST"])
def connections_page(request):
    if request.method == "POST":
        try:
            kind = request.POST["kind"]
            if kind in {"sheets", "google_sheets"} and request.POST.get("sheet_auth_mode"):
                from .google_oauth import sheet_config
                config = sheet_config(request.POST)
            else:
                config = json.loads(request.POST.get("config", "{}"))
            create_connection(request.user, request.POST.get("name", "데이터"), kind,
                              config, request.FILES.get("file"))
            messages.success(request, "데이터 연결을 등록했습니다. 연결 테스트로 확인해 주세요.")
            return redirect("connections")
        except PermissionDenied:
            raise
        except (ValueError, KeyError, TypeError):
            messages.error(request, "연결 설정 JSON, 파일 형식 및 크기를 확인해 주세요.")
    connections = list(connections_for(request.user))
    for connection in connections:
        editable_config = {key: value for key, value in connection.config.items() if key != "oauth_attempt"}
        connection.editable_config = json.dumps(editable_config, ensure_ascii=False, indent=2)
        connection.secret_names = sorted(connection.get_secrets())
        connection.file_extensions = sorted(FILE_EXTENSIONS.get(connection.kind, []))
    return render(request, "reportbuilder/connections.html", {"connections": connections,
                  "connectors": connector_catalog()})


@login_required
@require_POST
def connection_update(request, connection_id):
    if not request.user.is_staff:
        raise PermissionDenied("데이터 연결 수정은 관리자만 할 수 있습니다.")
    connection = get_object_or_404(Connection, pk=connection_id, owner=request.user)
    try:
        name = request.POST.get("name", "").strip()
        config = json.loads(request.POST.get("config", "{}"))
        secret_updates = json.loads(request.POST.get("secret_updates", "{}") or "{}")
        if (not name or not isinstance(config, dict) or not isinstance(secret_updates, dict)
                or len(json.dumps(config)) > 100000 or set(config) & (SECRET_FIELDS | {"path", "allowed_hosts", "url", "connection_string"})
                or "oauth_attempt" in config
                or set(secret_updates) - EDITABLE_SECRET_FIELDS
                or len(json.dumps(secret_updates)) > 100000
                or request.POST.get("clear_secrets") == "yes" and secret_updates):
            raise ValueError("연결 설정을 확인하세요.")
        if any((key in {"headers", "service_account"} and not isinstance(value, dict))
               or (key not in {"headers", "service_account"} and not isinstance(value, str))
               for key, value in secret_updates.items()):
            raise ValueError("비밀 정보의 형식을 확인하세요.")
        upload = request.FILES.get("file")
        if upload:
            allowed = FILE_EXTENSIONS.get(connection.kind)
            if not allowed or upload.size > 50 * 1024 * 1024 or Path(upload.name).suffix.lower() not in allowed:
                raise ValueError("파일 형식 또는 크기를 확인하세요.")
        secrets = {} if request.POST.get("clear_secrets") == "yes" else connection.get_secrets()
        secrets.update({key: value for key, value in secret_updates.items() if value not in (None, "")})
        if config.get("auth_mode") != "oauth" and connection.config.get("auth_mode") == "oauth":
            for key in ("oauth_client_id", "refresh_token", "access_token", "expires_at"):
                secrets.pop(key, None)
        old_upload = connection.upload if upload else None
        connection.name = name[:200]
        connection.config = config
        connection.set_secrets(secrets)
        if upload:
            connection.upload = upload
        connection.status = "UNTESTED"
        connection.last_test_at = None
        connection.save()
        if old_upload:
            old_upload.delete(save=False)
        audit(request.user, "connection_update", connection.pk, connector=connection.kind)
        messages.success(request, "데이터 연결을 수정했습니다. 연결 테스트로 확인해 주세요.")
    except (ValueError, TypeError, json.JSONDecodeError):
        messages.error(request, "연결 설정 JSON, 비밀 정보 또는 파일 형식을 확인해 주세요.")
    return redirect("connections")


@login_required
@require_POST
def connection_delete(request, connection_id):
    if not request.user.is_staff:
        raise PermissionDenied("데이터 연결 삭제는 관리자만 할 수 있습니다.")
    connection = get_object_or_404(Connection, pk=connection_id, owner=request.user)
    upload_name = connection.upload.name if connection.upload else None
    upload_storage = connection.upload.storage if upload_name else None
    audit(request.user, "connection_delete", connection.pk, connector=connection.kind)
    connection.delete()
    if upload_storage:
        upload_storage.delete(upload_name)
    messages.success(request, "데이터 연결을 삭제했습니다. 이 연결을 사용하는 보고서는 다시 매핑해야 합니다.")
    return redirect("connections")


@api
@require_http_methods(["GET", "POST"])
def connections_api(request):
    if request.method == "POST":
        value = body(request)
        connection = create_connection(request.user, value["name"], value["kind"], value.get("config", {}))
        return JsonResponse({"id": str(connection.pk), "name": connection.name, "kind": connection.kind}, status=201)
    return JsonResponse({"connections": [{"id": str(c.pk), "name": c.name, "kind": c.kind, "status": c.status}
                                        for c in connections_for(request.user)]})


@api
@require_POST
def connection_test_api(request, connection_id):
    connection = get_object_or_404(connections_for(request.user), pk=connection_id)
    result = test_connection(connection.kind, connection.runtime_config())
    connection.status = "CONNECTED"
    connection.last_test_at = timezone.now()
    connection.save(update_fields=["status", "last_test_at"])
    audit(request.user, "connection_test", connection.pk)
    return JsonResponse({"status": "CONNECTED", "result": result})


@api
def schema_api(request, connection_id):
    connection = get_object_or_404(connections_for(request.user), pk=connection_id)
    catalog = introspect(connection.kind, connection.runtime_config())
    if connection.allowed_objects:
        catalog["objects"] = [o for o in catalog["objects"] if o["object"] in connection.allowed_objects]
    for obj in catalog["objects"]:
        allowed = connection.allowed_columns.get(obj["object"])
        if allowed is not None:
            obj["columns"] = [c for c in obj["columns"] if c["column"] in allowed]
    return JsonResponse(catalog)


@api
@require_POST
def asset_api(request):
    upload = request.FILES.get("file") or request.FILES.get("image")
    if not upload or upload.size > 10 * 1024 * 1024:
        raise ValueError("Image size")
    content = normalised_png(upload)
    project = None
    if request.POST.get("report_id"):
        report = get_object_or_404(reports_for(request.user), pk=request.POST["report_id"])
        editable(report, request.user)
        project = report.project
    asset = Asset.objects.create(owner=request.user, name=Path(upload.name).name[:200], project=project,
                                 mime="image/png", file=ContentFile(content, name="image.png"))
    return JsonResponse({"id": str(asset.pk), "url": reverse("asset", args=[asset.pk])}, status=201)


def normalised_png(source):
    """Bound the decoded image and the re-encoded asset before persisting any file."""
    try:
        with Image.open(source) as image:
            if image.format not in {"PNG", "JPEG", "GIF", "WEBP"} or image.width * image.height > 20_000_000:
                raise ValueError("Unsupported image format or dimensions")
            image.load()
            output = io.BytesIO()
            image.convert("RGBA").save(output, format="PNG")
    except (OSError, Image.DecompressionBombError):
        raise ValueError("Invalid image") from None
    content = output.getvalue()
    validate_asset(content, "image/png")
    return content


@login_required
def asset_view(request, asset_id):
    asset = get_object_or_404(Asset, pk=asset_id)
    if not request.user.is_staff and asset.owner_id != request.user.pk:
        if not asset.project_id or not reports_for(request.user).filter(project_id=asset.project_id).exists():
            raise PermissionDenied
    return FileResponse(asset.file.open("rb"), content_type=asset.mime)


def project_assets(report):
    ids = set()
    for page in report.definition.get("pages", []):
        elements = list(page.get("elements", []))
        for band in page.get("bands", []):
            elements.extend(band.get("elements", []))
        ids.update(e["asset_id"] for e in elements if e.get("type") == "image" and e.get("asset_id"))
    result = {}
    for asset_id in ids:
        asset = get_object_or_404(Asset, pk=asset_id)
        if asset.owner_id != report.owner_id and asset.project_id != report.project_id:
            raise PermissionDenied
        with asset.file.open("rb") as stream:
            content = stream.read(MAX_ASSET + 1)
        validate_asset(content, asset.mime)
        result[str(asset.pk)] = {"content": content, "filename": "image.png", "mime": asset.mime}
    return result


@api
@require_POST
def project_import(request):
    upload = request.FILES.get("file") or request.FILES.get("project")
    if not upload or upload.size > 50 * 1024 * 1024:
        raise ValueError("Project size")
    value = import_project(upload.read())
    normalised = {asset_id: normalised_png(io.BytesIO(resource["content"]))
                  for asset_id, resource in value.get("assets", {}).items()}
    with transaction.atomic():
        project = Project.objects.create(name=str(value["name"])[:200], owner=request.user)
        asset_map = {}
        for old_id, content in normalised.items():
            asset = Asset.objects.create(owner=request.user, project=project, name="Imported image", mime="image/png",
                                         file=ContentFile(content, name="image.png"))
            asset_map[old_id] = str(asset.pk)
        reports = []
        for item in value["reports"]:
            definition = copy.deepcopy(item["definition"])
            for page in definition["pages"]:
                elements = list(page.get("elements", []))
                for band in page.get("bands", []):
                    elements.extend(band.get("elements", []))
                for element in elements:
                    if element.get("type") == "image":
                        element["asset_id"] = asset_map[element["asset_id"]]
            report = Report.objects.create(name=str(item["name"])[:200], project=project, owner=request.user,
                                           definition=definition, bindings=[])
            new_revision(report)
            reports.append(report)
    audit(request.user, "project_import", project.pk)
    if request.headers.get("Accept") == "application/json":
        return JsonResponse({"project_id": str(project.pk), "reports": [report_payload(r) for r in reports]})
    return redirect("designer", report_id=reports[0].pk)


def download(content, content_type, filename):
    response = HttpResponse(content, content_type=content_type)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    response["Cache-Control"] = "private, no-store"
    return response


def safe_cell(value):
    if value is None:
        return ""
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value if isinstance(value, (str, int, float, bool)) else str(value)


@api
def report_export(request, report_id, format):
    report = get_object_or_404(reports_for(request.user), pk=report_id)
    if format == "project":
        editable(report, request.user)
        content = export_project(report.project.name, [{"name": report.name, "definition": report.definition}],
                                 assets=project_assets(report))
        audit(request.user, "project_export", report.pk)
        return download(content, "application/zip", "report.wrpx")
    if format not in {"pdf", "html", "xlsx", "csv"}:
        raise ValueError("Export format")
    if request.GET.get("execution"):
        execution = get_object_or_404(Execution, pk=request.GET["execution"], report=report, user=request.user, status="SUCCEEDED")
        if execution.created_at < timezone.now() - timedelta(hours=24):
            raise PermissionDenied("결과 보존 기한이 지났습니다. 다시 실행해 주세요.")
        if execution.policy_fingerprint != policy_fingerprint(request.user, report, execution.bindings_snapshot):
            raise PermissionDenied("데이터 정책이 변경되었습니다. 다시 실행해 주세요.")
    else:
        execution, _ = execute_report(report, request.user, published=not (request.user.is_staff or report.owner_id == request.user.pk))
    audit(request.user, "result_export", report.pk, format=format, execution_id=str(execution.pk))
    if format == "pdf":
        return download(pdf_bytes(execution.rendered_html), "application/pdf", "report.pdf")
    if format == "html":
        return download(execution.rendered_html, "text/html; charset=utf-8", "report.html")
    data = execution.datasets_snapshot
    if format == "xlsx":
        book = Workbook()
        book.remove(book.active)
        for index, (dataset_id, result) in enumerate(data.items()):
            sheet = book.create_sheet(f"Data {index + 1}")
            fields = result.get("fields", [])
            ids = [f["field_id"] for f in fields] or list(result["rows"][0]) if result["rows"] else [f["field_id"] for f in fields]
            labels = {f["field_id"]: f.get("label", f.get("alias", f["field_id"])) for f in fields}
            sheet.append([safe_cell(labels.get(key, key)) for key in ids])
            for row in result["rows"]:
                sheet.append([safe_cell(row.get(key)) for key in ids])
            sheet.freeze_panes = "A2"
        if not book.sheetnames:
            book.create_sheet("Data")
        out = io.BytesIO()
        book.save(out)
        return download(out.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "report.xlsx")
    # Multiple datasets are separate CSV files; a single dataset remains a plain CSV.
    import zipfile
    csv_files = {}
    for index, (dataset_id, result) in enumerate(data.items()):
        fields = result.get("fields", [])
        ids = [f["field_id"] for f in fields] or (list(result["rows"][0]) if result["rows"] else [])
        text = io.StringIO()
        writer = csv.writer(text)
        writer.writerow([safe_cell(f.get("label", f.get("alias", f["field_id"]))) for f in fields] or ids)
        writer.writerows([[safe_cell(row.get(key)) for key in ids] for row in result["rows"]])
        csv_files[f"data-{index + 1}.csv"] = text.getvalue().encode("utf-8-sig")
    if len(csv_files) <= 1:
        return download(next(iter(csv_files.values()), b""), "text/csv; charset=utf-8", "report.csv")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in csv_files.items():
            archive.writestr(name, content)
    return download(out.getvalue(), "application/zip", "report-csv.zip")


@login_required
def publication_view(request, publication_id):
    publication = get_object_or_404(Publication, pk=publication_id, enabled=True, report__enabled=True)
    report = get_object_or_404(reports_for(request.user), pk=publication.report_id)
    response = render_viewer(request, report, published=True)
    if request.method == "GET":
        record_report_access(request, report, "view")
    return response


@api
@require_POST
def embed_session(request):
    value = body(request)
    report = get_object_or_404(reports_for(request.user), pk=value["report_id"])
    if not report.enabled or not hasattr(report, "publication") or not report.publication.enabled:
        raise ValueError("Not published")
    origin = value.get("origin")
    if origin not in settings.EMBED_ALLOWED_ORIGINS:
        raise PermissionDenied("임베드 호스트가 허용되지 않았습니다.")
    token = signing.dumps({"user_id": request.user.pk, "report_id": str(report.pk), "origin": origin,
                          "parameters": value.get("parameters", {}), "nonce": uuid.uuid4().hex}, salt="webreport-embed")
    return JsonResponse({"token": token, "url": reverse("embed"), "expires_in": 60})


@xframe_options_exempt
@csrf_exempt
@require_POST
def embed_view(request):
    # A server-issued short-lived signed token is submitted by host form POST. No third-party cookie required.
    from django.contrib.auth import get_user_model
    User = get_user_model()
    try:
        payload = signing.loads(request.POST.get("token", ""), salt="webreport-embed", max_age=60)
        if payload["origin"] not in settings.EMBED_ALLOWED_ORIGINS or request.headers.get("Origin") != payload["origin"]:
            raise PermissionDenied
        try:
            with transaction.atomic():
                EmbedNonce.objects.create(nonce=payload["nonce"], expires_at=timezone.now() + timedelta(minutes=2))
        except IntegrityError:
            raise PermissionDenied
        user = User.objects.get(pk=payload["user_id"], is_active=True)
        report = reports_for(user).get(pk=payload["report_id"])
        _, rendered = execute_report(report, user, payload.get("parameters", {}), published=True)
    except (signing.BadSignature, PermissionDenied, KeyError, ValueError, User.DoesNotExist, Report.DoesNotExist, DataError, DefinitionError):
        return HttpResponse("임베드 인증에 실패했습니다.", status=403)
    record_report_access(request, report, "embed")
    response = HttpResponse(rendered["html"])
    response["Content-Security-Policy"] = f"default-src 'none'; img-src data:; style-src 'unsafe-inline'; frame-ancestors {payload['origin']}"
    response["Cache-Control"] = "no-store"
    return response


@login_required
def manual(request):
    text = (Path(__file__).parent / "manual.md").read_text(encoding="utf-8")
    return render(request, "reportbuilder/manual.html", {"manual_html": markdown.markdown(text, extensions=["tables", "fenced_code"])})
