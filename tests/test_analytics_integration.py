import io
from datetime import timedelta

import pytest
from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.utils import timezone

from reportbuilder.definition import default_definition
from reportbuilder.models import Project, Publication, Report, ReportAccess, Revision

pytestmark = pytest.mark.django_db


@pytest.fixture
def analytics_report(client):
    user = get_user_model().objects.create_user(username="access-owner", is_staff=True)
    project = Project.objects.create(owner=user, name="Access report")
    report = Report.objects.create(owner=user, project=project, name="Access report", definition=default_definition(), revision=1)
    revision = Revision.objects.create(report=report, number=1, definition=report.definition, bindings=[])
    Publication.objects.create(report=report, revision=revision)
    client.force_login(user)
    return report


def test_successful_route_events_and_no_designer_or_statistics_self_count(client, analytics_report):
    report = analytics_report
    assert client.get(f"/reports/{report.pk}/").status_code == 200
    assert client.get(f"/published/{report.publication.pk}/").status_code == 200
    assert client.post(f"/api/reports/{report.pk}/execute/", data="{}", content_type="application/json").status_code == 200
    assert list(ReportAccess.objects.order_by("created_at").values_list("event", flat=True)) == ["view", "view", "execute"]
    assert client.head(f"/reports/{report.pk}/").status_code == 200
    assert client.get(f"/reports/{report.pk}/design/").status_code == 200
    assert client.post(f"/api/reports/{report.pk}/preview/", data="{}", content_type="application/json").status_code == 200
    assert client.get(f"/reports/{report.pk}/statistics/").status_code == 200
    assert client.get(f"/reports/{report.pk}/statistics/export/xlsx/").status_code == 200
    assert ReportAccess.objects.count() == 3
    assert client.post(f"/api/reports/{report.pk}/execute/", data='{"parameters":[]}', content_type="application/json").status_code == 400
    client.logout()
    assert client.get(f"/reports/{report.pk}/").status_code == 302
    assert client.post(f"/api/reports/{report.pk}/execute/", data="{}", content_type="application/json").status_code == 401
    assert ReportAccess.objects.count() == 3


def test_embed_records_once_and_disabled_recording_is_respected(client, analytics_report, settings):
    report = analytics_report
    settings.EMBED_ALLOWED_ORIGINS = ["https://example.test"]
    response = client.post("/api/embed-sessions/", data={"report_id": str(report.pk), "origin": "https://example.test"}, content_type="application/json")
    assert response.status_code == 200
    client.logout()
    token = response.json()["token"]
    assert client.post("/embed/", {"token": token}, HTTP_ORIGIN="https://example.test").status_code == 200
    assert client.post("/embed/", {"token": token}, HTTP_ORIGIN="https://example.test").status_code == 403
    assert ReportAccess.objects.get(report=report).event == "embed"
    settings.REPORT_ANALYTICS_ENABLED = False
    client.force_login(report.owner)
    assert client.get(f"/reports/{report.pk}/").status_code == 200
    assert ReportAccess.objects.count() == 1


def test_retention_removes_only_old_access_records(analytics_report):
    report = analytics_report
    old = ReportAccess.objects.create(report=report, event="view", created_at=timezone.now()-timedelta(days=366))
    recent = ReportAccess.objects.create(report=report, event="view")
    call_command("purge_report_access", days=365, stdout=io.StringIO())
    assert not ReportAccess.objects.filter(pk=old.pk).exists()
    assert ReportAccess.objects.filter(pk=recent.pk).exists()
    assert Report.objects.filter(pk=report.pk).exists()


def test_statistics_permissions_filters_and_export(client, analytics_report):
    from datetime import datetime, timezone as utc
    from openpyxl import load_workbook
    report = analytics_report
    report.name = '=HYPERLINK("https://example.test")'
    report.save(update_fields=["name"])
    ReportAccess.objects.create(report=report, event="view", country="KR", device="mobile", browser="safari", created_at=datetime(2026, 10, 1, 15, 0, tzinfo=utc.utc))
    ReportAccess.objects.create(report=report, event="execute", country="US", created_at=datetime(2026, 10, 1, 14, 59, tzinfo=utc.utc))
    url = f"/reports/{report.pk}/statistics/"
    query = {"start": "2026-10-02", "end": "2026-10-02", "country": "KR"}
    response = client.get(url, query)
    assert response.status_code == 200
    data = response.json()
    assert data["total"] == 1
    assert data["daily"] == [{"date": "2026-10-02", "count": 1}]
    assert data["recent"][0]["timestamp"].endswith("+09:00")
    export = client.get(url + "export/xlsx/", query)
    book = load_workbook(io.BytesIO(export.content))
    assert book["Summary"]["B2"].value.startswith("'=")
    assert book["Details"].max_row == 2
    assert client.get(url, {"start": "2026-10-03", "end": "2026-10-02"}).status_code == 400
    assert client.get(url, HTTP_AUTHORIZATION="Bearer invalid").status_code == 403
    stranger = get_user_model().objects.create_user(username="stats-stranger")
    client.force_login(stranger)
    assert client.get(url).status_code == 404
    assert client.get(url + "export/xlsx/").status_code == 404
    client.logout()
    assert client.get(url).status_code == 401


def test_proxy_trust_and_device_parsing(settings):
    from django.test import RequestFactory
    from reportbuilder.analytics import client_address, country_for, user_agent_dimensions
    factory = RequestFactory()
    settings.REPORT_TRUSTED_PROXIES = []
    request = factory.get("/", REMOTE_ADDR="10.0.0.1", HTTP_X_FORWARDED_FOR="8.8.8.8")
    assert str(client_address(request)) == "10.0.0.1"
    assert country_for(client_address(request)) == "Private"
    settings.REPORT_TRUSTED_PROXIES = ["10.0.0.0/8"]
    assert str(client_address(request)) == "8.8.8.8"
    request.META["HTTP_X_FORWARDED_FOR"] = "1.1.1.1, 8.8.8.8"
    assert str(client_address(request)) == "8.8.8.8"
    dimensions = user_agent_dimensions("Mozilla/5.0 (iPhone; CPU iPhone OS 18_0 like Mac OS X) Mobile/15E148 Safari/604.1")
    assert dimensions == {"device": "mobile", "browser": "safari", "os": "ios"}
