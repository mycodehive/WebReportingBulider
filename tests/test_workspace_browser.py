"""Real browser integration; CI installs Chromium and uploads synthetic screenshots."""
from pathlib import Path

import pytest
from django.contrib.auth import get_user_model
from django.core.management import call_command

from reportbuilder.models import Report, ReportAccess


@pytest.mark.django_db(transaction=True)
def test_workspace_theme_fullscreen_and_statistics_downloads(live_server, client, settings, tmp_path):
    playwright = pytest.importorskip("playwright.sync_api")
    settings.MEDIA_ROOT = tmp_path / "media"
    user = get_user_model().objects.create_user(username="visual-test", is_staff=True)
    call_command("seed_demo", username=user.username, verbosity=0)
    report = Report.objects.get(owner=user, name="매출 현황 예제")
    client.force_login(user)
    ReportAccess.objects.create(report=report, event="view", country="KR", device="mobile", browser="safari", os="ios")
    with playwright.sync_playwright() as api:
        try:
            browser = api.chromium.launch()
        except playwright.Error as exc:
            if "Executable doesn't exist" in str(exc):
                pytest.skip("Chromium binary is unavailable")
            raise
        context = browser.new_context(viewport={"width": 1440, "height": 1000}, locale="ko-KR")
        context.add_cookies([{"name": settings.SESSION_COOKIE_NAME,
                             "value": client.cookies[settings.SESSION_COOKIE_NAME].value,
                             "url": live_server.url}])
        page = context.new_page()
        errors = []
        page.on("pageerror", lambda error: errors.append(str(error)))
        shots = Path("test-results")
        shots.mkdir(exist_ok=True)
        try:
            page.goto(f"{live_server.url}/reports/{report.pk}/design/")
            playwright.expect(page.locator("#report-name")).to_have_value("매출 현황 예제")
            playwright.expect(page.locator("#report-canvas .canvas-element").first).to_be_visible()
            page.screenshot(path=str(shots / "workspace-light.png"), full_page=True)
            old_width = page.locator("#designer-workspace").bounding_box()["width"]
            page.locator("#sidebar-toggle").click()
            playwright.expect(page.locator("html")).to_have_attribute("data-sidebar", "collapsed")
            page.wait_for_function("old => document.querySelector('#designer-workspace').getBoundingClientRect().width > old", arg=old_width)
            page.locator("#theme-toggle").click()
            playwright.expect(page.locator("html")).to_have_attribute("data-theme", "dark")
            page.reload()
            playwright.expect(page.locator("html")).to_have_attribute("data-theme", "dark")
            playwright.expect(page.locator("html")).to_have_attribute("data-sidebar", "collapsed")
            playwright.expect(page.locator("#report-canvas .canvas-element").first).to_be_visible()
            colors = page.locator("#report-name").evaluate("e => ({fg:getComputedStyle(e).color,bg:getComputedStyle(e).backgroundColor})")
            assert colors["fg"] != colors["bg"]
            assert page.locator("#report-canvas").evaluate("e => getComputedStyle(e).backgroundColor") == "rgb(255, 255, 255)"
            page.locator("#workspace-fullscreen").click()
            playwright.expect(page.locator("#workspace-fullscreen")).to_have_attribute("aria-pressed", "true")
            page.wait_for_function("document.querySelector('#designer').getBoundingClientRect().width >= innerWidth - 2")
            playwright.expect(page.locator("#save")).to_be_visible()
            page.screenshot(path=str(shots / "workspace-dark-fullscreen.png"))
            page.locator("#preview").click()
            playwright.expect(page.locator("#preview-dialog")).to_be_visible()
            page.locator('[data-close="preview-dialog"]').click()
            page.locator("#workspace-fullscreen").click()
            playwright.expect(page.locator("#workspace-fullscreen")).to_have_attribute("aria-pressed", "false")
            page.goto(f"{live_server.url}/reports/{report.pk}/")
            page.goto(f"{live_server.url}/reports/")
            page.locator(f'[data-report-statistics="{report.pk}"]').click()
            playwright.expect(page.locator("#statistics-dialog")).to_be_visible()
            playwright.expect(page.locator('[data-stat-kpi="total"]')).to_contain_text("2")
            page.screenshot(path=str(shots / "statistics-dark.png"))
            with page.expect_download() as downloaded:
                page.locator('[data-statistics-export="xlsx"]').click()
            from openpyxl import load_workbook
            workbook = load_workbook(downloaded.value.path())
            assert len(workbook.sheetnames) >= 4
            with page.expect_download(timeout=60000) as downloaded:
                page.locator('[data-statistics-export="pdf"]').click()
            from pypdf import PdfReader
            assert len(PdfReader(downloaded.value.path()).pages) >= 1
            page.keyboard.press("Escape")
            playwright.expect(page.locator("#statistics-dialog")).not_to_be_visible()
            page.locator("#theme-toggle").click()
            page.locator(f'[data-report-statistics="{report.pk}"]').click()
            playwright.expect(page.locator('[data-stat-kpi="total"]')).to_contain_text("2")
            page.screenshot(path=str(shots / "statistics-light.png"))
            assert errors == []
        except Exception:
            page.screenshot(path=str(shots / "browser-failure.png"), full_page=True)
            raise
        finally:
            context.close()
            browser.close()
